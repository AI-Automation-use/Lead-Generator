# from docx import Document
# from io import BytesIO

# def create_lead_docx(lead_analysis: str, company: str):
#     doc = Document()
#     doc.add_heading(f"Lead Analysis for {company}", 0)
#     doc.add_paragraph(lead_analysis)
#     buf = BytesIO()
#     doc.save(buf)
#     buf.seek(0)
#     return f"{company}_lead_analysis.docx", buf


import logging
import requests
import os
import base64
from typing import List, Optional, Tuple
from io import BytesIO
from docx import Document
from auth import _get_app_token, GRAPH_API_ENDPOINT
from utils import _addr_list

def create_lead_docx(lead_analysis: str, company: str) -> Tuple[str, BytesIO]:
    doc = Document()
    doc.add_heading(f"Lead Analysis for {company}", 0)
    doc.add_paragraph(lead_analysis)
    doc_stream = BytesIO()
    doc.save(doc_stream)
    doc_stream.seek(0)
    filename = f"{company}_lead_analysis.docx"
    return filename, doc_stream


def create_full_docx(website_content: str, linkedin_content: str, news_content: str, company: str) -> Tuple[str, BytesIO]:
    doc = Document()
    doc.add_heading("Company Content and Analysis", 0)
    doc.add_heading("Company Website Content:", level=1)
    doc.add_paragraph(website_content)
    doc.add_heading("Company LinkedIn Profile Content:", level=1)
    doc.add_paragraph(linkedin_content)
    doc.add_heading("Recent News Articles:", level=1)
    doc.add_paragraph(news_content)
    doc_stream = BytesIO()
    doc.save(doc_stream)
    doc_stream.seek(0)
    filename = f"{company}_full_content.docx"
    return filename, doc_stream

# -----------------------------
# External API / Email helpers
# -----------------------------

def send_email_app_only(
    sender_user_id: str,
    to_emails: List[str],
    subject: str,
    html_body: str,
    cc_emails: Optional[List[str]] = None,
    bcc_emails: Optional[List[str]] = None,
    attachments: Optional[List[tuple]] = None,
) -> bool:
    try:
        token = _get_app_token()
        message = {
            "message": {
                "subject": subject,
                "importance": "Normal",
                "body": {"contentType": "HTML", "content": html_body},
                "toRecipients": _addr_list(to_emails),
                "attachments": [],
            },
            "saveToSentItems": "true",
        }
        if cc_emails:
            message["message"]["ccRecipients"] = _addr_list(cc_emails)
        if bcc_emails:
            message["message"]["bccRecipients"] = _addr_list(bcc_emails)
        if attachments:
            for name, data_stream in attachments:
                data_stream.seek(0)
                encoded_content = base64.b64encode(data_stream.read()).decode("utf-8")
                message["message"]["attachments"].append({
                    "@odata.type": "#microsoft.graph.fileAttachment",
                    "name": name,
                    "contentType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    "contentBytes": encoded_content,
                })
        url = f"{GRAPH_API_ENDPOINT}/users/{sender_user_id}/sendMail"
        r = requests.post(url, headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"}, json=message, timeout=30)
        if r.status_code == 202:
            logging.info(f"✅ Email Sent Successfully to {to_emails}.")
            return True
        else:
            logging.error(f"❌ Send failed [{r.status_code}]: {r.text}")
            return False
    except Exception as e:
        logging.error(f"❌ Error sending email: {str(e)}")
        return False


def send_lead_data_to_api(
    lead_areas: str,
    account_name: str,
    lead_name: str,
    file_name: Optional[str] = None,
    file_bytes: Optional[bytes] = None,
) -> bool:
    api_url = os.getenv("LEAD_API_URL")
    if not api_url:
        logging.error("Error: 'LEAD_API_URL' environment variable not found.")
        return False
    try:
        if file_name and file_bytes is not None:
            logging.info("Attempting to send lead data + file to API (multipart/form-data)...")
            files = [
                ('new_leadidentificationarea', (None, lead_areas)),
                ('new_name', (None, lead_name)),
                ('new_accountname', (None, account_name)),
                ('new_supportingdocuments', (file_name, BytesIO(file_bytes), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')),
            ]
            resp = requests.post(api_url, files=files, timeout=60)
        else:
            logging.info("Attempting to send lead data to API (JSON)...")
            data = {"new_leadidentificationarea": lead_areas, "new_name": lead_name, "new_accountname": account_name}
            resp = requests.post(api_url, json=data, timeout=60)
        resp.raise_for_status()
        logging.info("✅ Lead data successfully sent to API.")
        return True
    except requests.exceptions.RequestException as e:
        logging.error(f"❌ Failed to send lead data to API: {e}")
        if getattr(e, "response", None) is not None:
            logging.error(f"Status: {e.response.status_code}")
            logging.error(f"Response: {e.response.text}")
        return False
